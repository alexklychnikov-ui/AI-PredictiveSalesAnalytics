"""Generate Stage 10 TZ as DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x59, 0x59, 0x59)
OUT = Path("docs") / "TZ_AI_prediktivnaya_analitika_prodazh.docx"


def set_run_font(run, size=11, bold=False, color=None, name="Calibri") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_horizontal_line(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F4E79")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def h(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), size=14, bold=True, color=ACCENT)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)


def body(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), size=11)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    set_run_font(p.add_run(text), size=11)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(2)


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("Техническое задание"), size=14, bold=True, color=ACCENT)

    main_title = doc.add_paragraph()
    main_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        main_title.add_run("AI-система предиктивной аналитики продаж"),
        size=22,
        bold=True,
        color=ACCENT,
    )
    add_horizontal_line(main_title)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        meta.add_run("MVP · веб-приложение · Streamlit + PostgreSQL + Prophet"),
        size=10,
        color=MUTED,
    )
    doc.add_paragraph()

    h(doc, "1. Название и описание проекта")
    body(
        doc,
        "Продукт — компактное веб-приложение для загрузки истории продаж, "
        "проверки качества данных, построения проверяемого прогноза, сравнения сценариев "
        "и выдачи рекомендаций с явными основаниями и ограничениями.",
    )
    body(
        doc,
        "Система универсальна для розницы, e-commerce, услуг и B2B за счёт сопоставления "
        "произвольных колонок с внутренней схемой и выбора KPI / частоты агрегации. "
        "Интерфейс MVP — на русском. AI (OpenAI) опционален: без ключа работают аналитика, "
        "прогноз, сценарии и шаблонные рекомендации.",
    )

    h(doc, "2. Целевая аудитория и боли")
    body(doc, "Владелец малого / среднего бизнеса")
    bullet(doc, "не понимает ожидаемую выручку на следующий период;")
    bullet(doc, "планирует закупки и расходы интуитивно;")
    bullet(doc, "нет отдельной команды аналитиков.")

    body(doc, "Руководитель продаж")
    bullet(doc, "отчёты собираются вручную;")
    bullet(doc, "поздно замечает падения и аномалии;")
    bullet(doc, "сложно быстро сравнить базовый и стрессовый сценарии.")

    body(doc, "Маркетолог")
    bullet(doc, "неясно, в какие периоды запускать акции;")
    bullet(doc, "слабо видна связь скидок / рекламы с продажами;")
    bullet(doc, "риск принять корреляцию за причинность.")

    body(doc, "Аналитик")
    bullet(doc, "много времени уходит на очистку и приведение данных;")
    bullet(doc, "нет единого воспроизводимого backtesting;")
    bullet(doc, "метрики, графики и выводы собираются вручную.")

    h(doc, "3. Функциональные требования")
    bullet(doc, "Загрузка CSV/XLSX, предпросмотр, сопоставление колонок, выбор KPI и частоты.")
    bullet(doc, "Контроль качества данных и сохранение нормализованного ряда в PostgreSQL.")
    bullet(doc, "Аналитика: KPI, тренд, сезонность, аномалии, корреляции.")
    bullet(
        doc,
        "Прогноз (Prophet + seasonal-naive), доверительный интервал, rolling backtesting, "
        "метрики MAE/RMSE/WAPE/sMAPE.",
    )
    bullet(
        doc,
        "Сценарии: базовый / оптимистичный / пессимистичный / пользовательский; "
        "what-if по пригодным факторам.",
    )
    bullet(
        doc,
        "Рекомендации: rule-based всегда; OpenAI — только по агрегированным фактам, "
        "с валидацией чисел.",
    )
    bullet(doc, "Экспорт прогноза/сценариев (CSV) и отчёта рекомендаций (TXT).")
    bullet(doc, "Демо-датасет без ПДн; развёртывание на VPS (Docker Compose) с HTTPS и Basic Auth.")

    h(doc, "4. Ограничения и риски")
    body(doc, "Ограничения")
    bullet(doc, "MVP — один пользовательский контур dashboard, без полноценного multi-tenant SaaS.")
    bullet(doc, "Точность не гарантируется на коротких / редких / сильно нестабильных рядах.")
    bullet(doc, "Стресс-% и корреляции факторов не равны причинному эффекту.")
    bullet(doc, "OpenAI зависит от внешнего API, ключа и бюджета токенов.")
    bullet(doc, "Бюджет проекта ограничен рамками учебного / портфолио MVP (VPS + опциональный OpenAI).")

    body(doc, "Риски")
    bullet(doc, "Нехватка или плохое качество входных данных → прогноз с низкой уверенностью или отказ.")
    bullet(doc, "Утечка секретов / ПДн при неаккуратной загрузке → лимиты, санитизация, без сырых ПДн в AI.")
    bullet(doc, "Ограничения VPS (CPU/RAM) на длинных рядах и Prophet.")
    bullet(doc, "Изменение API OpenAI или недоступность ключа → обязательный fallback на правила.")
    bullet(doc, "Переоценка AI-рекомендаций пользователем без учёта оговорок системы.")

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        footer.add_run("Документ фиксирует границы MVP на этапе сдачи проекта."),
        size=9,
        color=MUTED,
    )

    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
